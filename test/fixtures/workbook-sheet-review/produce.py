"""Independent openpyxl fixture producer. Run with bundled Python; convert with LO separately."""
import json, pathlib, openpyxl
from openpyxl.styles import PatternFill
P=pathlib.Path(__file__).parent
for changed in (False,True):
 w=openpyxl.Workbook();w.remove(w.active)
 names=['Overview','Hidden Δ','Secret','Empty','Café 中','Removed'] if not changed else ['Added','Café 中','Overview','Secret','Hidden Δ','Empty','Renamed']
 for name in names:
  s=w.create_sheet(name)
  if name!='Empty':
   s.append(['Key','Value']);s.append(['one',10]);s.append(['two',20])
  if name=='Overview': s['B4']='=SUM(B2:B3)'
  if name=='Hidden Δ': s.sheet_state='hidden';s['B3']=21 if changed else 20
  if name=='Secret': s.sheet_state='hidden' if changed else 'veryHidden'
  if name=='Café 中': s['A2']='Unicode Ω';s['B2'].number_format='0.000' if changed else '0.00'
  if name=='Empty': s['C5'].fill=PatternFill('solid',fgColor='FFFFCC')
 w.save(P/('changed.xlsx' if changed else 'baseline.xlsx'))
# Readback of the actual ZIP files by independent producer's reader (no app decoder).
report={}
for file in ['baseline.xlsx','changed.xlsx']:
 w=openpyxl.load_workbook(P/file,data_only=False)
 report[file]=[{'name':s.title,'ordinal':i,'visibility':s.sheet_state,'cells':{c.coordinate:{'value':c.value,'type':c.data_type,'numberFormat':c.number_format} for row in s for c in row if c.value is not None or c.has_style}} for i,s in enumerate(w)]
(P/'openpyxl-expected.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
